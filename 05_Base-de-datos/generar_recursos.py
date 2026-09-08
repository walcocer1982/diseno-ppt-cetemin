# -*- coding: utf-8 -*-
"""Recursos de aprendizaje autónomo — de Markdown a Word.

Son las lecturas que el estudiante hace en el bloque asincrónico ANTES de las
sesiones que preparan, y de las que salen las preguntas del Cuestionario de
Verificación. El bloque dura 135 min: 90 de lectura y 45 de cuestionario.

    CV1  antes de la sesión 1   prepara las sesiones 1 a 6
    CV2  después de la 6        prepara las sesiones 7 a 10
    CV3  después del EP         prepara las sesiones 13 a 18
    CV4  después de la 18       prepara las sesiones 19 a 22

El texto vive en `SI/recursos-autonomos/*.md` y no aquí: son miles de palabras,
y el instructor tiene que poder corregir una frase sin abrir un archivo de
código. Este script solo les pone la forma.

El dialecto de Markdown que entiende es corto a propósito:

    @clave: valor     metadatos de la portada (solo al inicio del archivo)
    ## texto          capítulo — empieza en página nueva
    ### texto         apartado
    - texto           viñeta
    | a | b |         tabla (la primera fila es la cabecera)
    > texto           recuadro «Lo que se pregunta»
    **texto**         negrita   ·   *texto*   cursiva

Uso:  python generar_recursos.py [CV1 CV2 ...]      (sin argumentos, todos)
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / "05_Base-de-datos" / "SI" / "recursos-autonomos"
DEST = (RAIZ / "03_Entregables-diseño" / "SI - Ciclo II"
        / "1 - SI-SGCSSMA - Sistema de gestión de calidad, seguridad, salud y medio ambiente"
        / "Instrumentos de evaluación" / "CUESTIONARIOS Y EXÁMENES")

AZUL = RGBColor(0x0D, 0x26, 0x32)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GRIS = RGBColor(0x5A, 0x6A, 0x72)
LETRA = "Aptos"

for _f in (sys.stdout, sys.stderr):
    try:
        _f.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


# ─────────────────────────────── el Markdown ───────────────────────────────

def leer(ruta: Path):
    """Devuelve (metadatos, bloques). Un bloque es (tipo, contenido)."""
    meta, bloques, i = {}, [], 0
    lineas = ruta.read_text(encoding="utf-8").splitlines()

    while i < len(lineas) and (not lineas[i].strip() or lineas[i].startswith("@")):
        if lineas[i].startswith("@"):
            k, _, v = lineas[i][1:].partition(":")
            meta[k.strip()] = v.strip()
        i += 1

    while i < len(lineas):
        ln = lineas[i]
        if not ln.strip():
            i += 1
        elif ln.startswith("### "):
            bloques.append(("apartado", ln[4:].strip())); i += 1
        elif ln.startswith("## "):
            bloques.append(("capitulo", ln[3:].strip())); i += 1
        elif ln.startswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].startswith("|"):
                celdas = [c.strip() for c in lineas[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in celdas):
                    filas.append(celdas)
                i += 1
            bloques.append(("tabla", filas))
        elif ln.startswith("> "):
            trozos = []
            while i < len(lineas) and lineas[i].startswith(">"):
                trozos.append(lineas[i].lstrip(">").strip()); i += 1
            bloques.append(("recuadro", [t for t in trozos if t]))
        elif ln.startswith("- "):
            trozos = []
            while i < len(lineas) and lineas[i].startswith("- "):
                trozos.append(lineas[i][2:].strip()); i += 1
            bloques.append(("vinetas", trozos))
        else:
            trozos = []
            while i < len(lineas) and lineas[i].strip() and not re.match(
                    r"^(#{2,3} |\||> |- )", lineas[i]):
                trozos.append(lineas[i].strip()); i += 1
            bloques.append(("parrafo", " ".join(trozos)))
    return meta, bloques


def escribir(p, texto, px=11, tinta=AZUL, negrita=False):
    """Vuelca el texto en el párrafo respetando **negrita** y *cursiva*."""
    for trozo in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
        if not trozo:
            continue
        neg, cur = negrita, False
        if trozo.startswith("**"):
            trozo, neg = trozo[2:-2], True
        elif trozo.startswith("*"):
            trozo, cur = trozo[1:-1], True
        r = p.add_run(trozo)
        r.font.name = LETRA
        r.font.size = Pt(px)
        r.font.bold = neg
        r.font.italic = cur
        r.font.color.rgb = tinta
    return p


def parrafo(doc, texto="", px=11, tinta=AZUL, negrita=False, antes=0, despues=6,
            izq=0, alineacion=None, interlinea=1.25):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(antes), Pt(despues)
    f.left_indent, f.line_spacing = Cm(izq), interlinea
    if alineacion is not None:
        p.alignment = alineacion
    return escribir(p, texto, px, tinta, negrita)


def sombrear(celda, color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    e = OxmlElement("w:shd")
    e.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(e)


# ─────────────────────────────── el documento ───────────────────────────────

def portada(doc, meta):
    parrafo(doc, meta.get("titulo", ""), 22, AZUL, True, despues=4, interlinea=1.1)
    parrafo(doc, meta.get("subtitulo", ""), 14, TEAL, despues=18, interlinea=1.1)

    ficha = [("Curso", "curso"), ("Prepara", "prepara"),
             ("Duración", "tiempo"), ("Cuándo se lee", "cuando")]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for etiqueta, clave in ficha:
        if not meta.get(clave):
            continue
        c = t.add_row().cells
        c[0].width, c[1].width = Cm(3.6), Cm(12.4)
        escribir(c[0].paragraphs[0], etiqueta, 10, GRIS, True)
        escribir(c[1].paragraphs[0], meta[clave], 10, AZUL)


def recuadro(doc, lineas):
    t = doc.add_table(rows=1, cols=1)
    celda = t.rows[0].cells[0]
    sombrear(celda, "FFF7D6")
    celda.paragraphs[0]._p.getparent().remove(celda.paragraphs[0]._p)
    for j, ln in enumerate(lineas):
        p = celda.add_paragraph()
        p.paragraph_format.space_before = Pt(6 if j == 0 else 0)
        p.paragraph_format.space_after = Pt(6 if j == len(lineas) - 1 else 2)
        p.paragraph_format.left_indent = Cm(0.3)
        escribir(p, ln, 10.5, AZUL)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def tabla(doc, filas):
    t = doc.add_table(rows=0, cols=len(filas[0]))
    t.style = "Table Grid"
    for n, fila in enumerate(filas):
        celdas = t.add_row().cells
        for celda, txt in zip(celdas, fila):
            if n == 0:
                sombrear(celda, "EEF2F4")
            p = celda.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            escribir(p, txt, 10, AZUL if n == 0 else GRIS, negrita=(n == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def construir(meta, bloques):
    doc = Document()
    est = doc.styles["Normal"]
    est.font.name, est.font.size = LETRA, Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)

    portada(doc, meta)
    primero = True
    for tipo, cont in bloques:
        if tipo == "capitulo":
            if not primero:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            primero = False
            parrafo(doc, cont, 16, TEAL, True, antes=10, despues=10, interlinea=1.1)
        elif tipo == "apartado":
            parrafo(doc, cont, 12.5, AZUL, True, antes=14, despues=6, interlinea=1.1)
        elif tipo == "parrafo":
            parrafo(doc, cont, 11, AZUL, alineacion=WD_ALIGN_PARAGRAPH.JUSTIFY)
        elif tipo == "vinetas":
            for v in cont:
                parrafo(doc, "•  " + v, 11, AZUL, izq=0.6, despues=3)
        elif tipo == "tabla":
            tabla(doc, cont)
        elif tipo == "recuadro":
            recuadro(doc, cont)
    return doc


def main(argv) -> int:
    pedidos = [a.upper() for a in argv[1:]]
    fuentes = sorted(FUENTE.glob("*.md"))
    if pedidos:
        fuentes = [f for f in fuentes if f.stem.upper() in pedidos]
    if not fuentes:
        print("No hay nada que generar en %s" % FUENTE)
        return 1

    print("Recursos de aprendizaje autónomo · SI-SGCSSMA")
    for md in fuentes:
        meta, bloques = leer(md)
        doc = construir(meta, bloques)
        caps = sum(1 for t, _ in bloques if t == "capitulo")
        palabras = sum(len(c.split()) for t, c in bloques if t == "parrafo")
        palabras += sum(len(" ".join(c).split()) for t, c in bloques
                        if t in ("vinetas", "recuadro"))
        nombre = "%s Recurso autónomo.docx" % md.stem.upper()
        carpeta = DEST / md.stem.upper()
        carpeta.mkdir(parents=True, exist_ok=True)
        try:
            doc.save(carpeta / nombre)
        except PermissionError:
            print("   %-4s NO se pudo escribir: %s está abierto" % (md.stem, nombre))
            continue
        # 120 palabras por minuto es lectura tecnica con comprension, no lectura
        # de novela: es la referencia con la que se calibran los 90 minutos.
        print("   %-4s %d capítulos · %d palabras · ~%d min de lectura · %s"
              % (md.stem.upper(), caps, palabras, round(palabras / 120), nombre))
    print("   destino: %s" % DEST)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
