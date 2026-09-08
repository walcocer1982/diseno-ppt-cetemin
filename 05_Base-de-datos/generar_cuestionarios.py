# -*- coding: utf-8 -*-
"""Bancos de preguntas de SI-SGCSSMA — CV1 a CV4, examen parcial y examen final.

Cada banco tiene **50 preguntas** de cuatro alternativas. Son BANCOS, no
exámenes fijos: el instructor arma cada aplicación tomando de aquí, o la
plataforma las sortea.

De dónde salen las preguntas:

    CV1  del cuadernillo de lectura autónoma 1   prepara las sesiones 1 a 6
    CV2  del cuadernillo de lectura autónoma 2   prepara las sesiones 7 a 10
    EP   de las clases                           sesiones 1 a 10
    CV3  del cuadernillo de lectura autónoma 3   prepara las sesiones 13 a 18
    CV4  del cuadernillo de lectura autónoma 4   prepara las sesiones 19 a 22
    EF   de las clases                           sesiones 13 a 22

El orden del curso de 96 h es: CV1 · sesiones 1-6 · CV2 · sesiones 7-12 · EP ·
CV3 · sesiones 13-18 · CV4 · sesiones 19-24 · EF. Los CV se responden ANTES de
las sesiones que preparan, así que toda pregunta de un CV tiene que poder
contestarse con el cuadernillo y con nada más.

**La alternativa correcta va resaltada en amarillo, en su sitio.** No hay hoja
de claves al final: quien pasa el banco a Blackboard copia la pregunta, las
cuatro alternativas y la retroalimentación sin moverse de la página. Por lo
mismo, el archivo NO se entrega ni se imprime para el estudiante.

REGLA DE CONTENIDO: ninguna pregunta afirma nada que no esté en el cuadernillo
o verificado en el §15. La línea «Fuente» de cada ítem dice de dónde sale la
clave, y es lo que permite revisar el banco dentro de un año sin volver a
discutirlo.

Las preguntas viven en `SI/preguntas/*.csv`, no aquí: son trescientas, y el
instructor tiene que poder corregir una sin abrir un archivo de código.

Uso:  python generar_cuestionarios.py [CV1 EP ...]      (sin argumentos, todos)
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / "05_Base-de-datos" / "SI" / "preguntas"
DEST = (RAIZ / "03_Entregables-diseño" / "SI - Ciclo II"
        / "1 - SI-SGCSSMA - Sistema de gestión de calidad, seguridad, salud y medio ambiente"
        / "Instrumentos de evaluación" / "CUESTIONARIOS Y EXÁMENES")

AZUL = RGBColor(0x0D, 0x26, 0x32)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
GRIS = RGBColor(0x5A, 0x6A, 0x72)
LETRA = "Aptos"

# El nombre largo de cada evento. La cobertura y el orden están arriba, en el
# docstring: aquí solo hace falta el rótulo que va en la portada.
BANCOS = {
    "CV1": "Cuestionario de Verificación 1",
    "CV2": "Cuestionario de Verificación 2",
    "EP":  "Examen Parcial",
    "CV3": "Cuestionario de Verificación 3",
    "CV4": "Cuestionario de Verificación 4",
    "EF":  "Examen Final",
}

for _f in (sys.stdout, sys.stderr):
    try:
        _f.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def equilibrar(items):
    """Reparte las claves entre a, b, c y d.

    Escritas de corrido, las claves se amontonan en b y c: un banco así se
    aprueba respondiendo siempre lo mismo. Aquí la alternativa correcta se mueve
    a la posición que le toca por turno, y las otras tres conservan su orden.
    Es determinista: el banco sale igual cada vez que se genera.
    """
    LETRAS = "abcd"
    salida = []
    for i, it in enumerate(items):
        ses, preg, a, b, c, d, clave, porque, fuente = it
        alts = [a, b, c, d]
        j = LETRAS.index(clave.strip().lower())
        resto = [x for k, x in enumerate(alts) if k != j]
        destino = i % 4
        nuevas = resto[:destino] + [alts[j]] + resto[destino:]
        salida.append((ses, preg, *nuevas, LETRAS[destino], porque, fuente))
    return salida


def escribir(p, texto, px=11, tinta=AZUL, negrita=False, resaltado=False):
    r = p.add_run(texto)
    r.font.name = LETRA
    r.font.size = Pt(px)
    r.font.bold = negrita
    r.font.color.rgb = tinta
    if resaltado:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def parrafo(doc, texto="", px=11, tinta=AZUL, negrita=False, antes=0, despues=4,
            izq=0.0, resaltado=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    p.paragraph_format.left_indent = Cm(izq)
    escribir(p, texto, px, tinta, negrita, resaltado)
    return p


def nota(doc, etiqueta, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.space_after = Pt(0)
    escribir(p, etiqueta, 9.5, TEAL, True)
    escribir(p, texto, 9.5, GRIS)


def documento(codigo, items):
    doc = Document()
    est = doc.styles["Normal"]
    est.font.name, est.font.size = LETRA, Pt(11)
    for s_ in doc.sections:
        s_.left_margin = s_.right_margin = Cm(2.2)

    parrafo(doc, "BANCO DE PREGUNTAS · %s" % codigo, 20, AZUL, True, despues=2)
    parrafo(doc, BANCOS[codigo], 14, TEAL, despues=14)

    parrafo(doc, "PREGUNTAS", 13, AZUL, True, antes=18, despues=6)
    for n, (ses, preg, a, b, c, d, clave, porque, fuente) in enumerate(items, 1):
        parrafo(doc, "%d.  %s" % (n, preg), 11, AZUL, True, antes=13, despues=4)
        for letra, alt in zip("abcd", (a, b, c, d)):
            parrafo(doc, "%s)  %s" % (letra, alt), 11,
                    AZUL if letra == clave else GRIS,
                    izq=0.7, despues=2, resaltado=(letra == clave))
        nota(doc, "Por qué:  ", porque)
        nota(doc, "Fuente:  ", "%s  ·  %s" % (fuente, ses))
    return doc


def main(argv) -> int:
    DEST.mkdir(parents=True, exist_ok=True)
    pedidos = [a.upper() for a in argv[1:]]
    codigos = [c for c in BANCOS if not pedidos or c in pedidos]
    if pedidos and not codigos:
        print("No conozco: %s" % ", ".join(pedidos))
        return 1

    print("Bancos de preguntas de SI-SGCSSMA")
    hubo = False
    for codigo in codigos:
        ruta = FUENTE / ("%s.csv" % codigo)
        if not ruta.exists():
            if pedidos:
                print("   %-4s todavía no tiene preguntas (falta %s)" % (codigo, ruta.name))
            continue
        hubo = True
        with ruta.open(encoding="utf-8-sig", newline="") as f:
            items = [tuple(fila) for fila in csv.reader(f)][1:]
        items = equilibrar(items)
        archivo = "%s Banco de preguntas.docx" % codigo
        # Los CV llevan carpeta propia, porque cada uno va con su cuadernillo.
        # El EP y el EF no tienen cuadernillo: van sueltos, al mismo nivel.
        carpeta = DEST / codigo if codigo.startswith("CV") else DEST
        carpeta.mkdir(parents=True, exist_ok=True)
        try:
            documento(codigo, items).save(carpeta / archivo)
        except PermissionError:
            print("   %-4s NO se pudo escribir: %s está abierto" % (codigo, archivo))
            continue
        claves, sesiones = {}, {}
        for it in items:
            claves[it[6]] = claves.get(it[6], 0) + 1
            sesiones[it[0]] = sesiones.get(it[0], 0) + 1
        aviso = "" if len(items) == 50 else "   <-- no son 50"
        print("   %-4s %d preguntas · claves %s · %s · %s%s"
              % (codigo, len(items), dict(sorted(claves.items())),
                 dict(sorted(sesiones.items())), archivo, aviso))
    if not hubo:
        print("   no hay ningún CSV de preguntas en %s" % FUENTE)
    else:
        print("   destino: %s" % DEST)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
